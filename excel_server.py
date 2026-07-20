from __future__ import annotations

import json
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
EXCEL_PATH = ROOT / "资源支持价值评估评分表2026版.xlsx"
RECORDS_PATH = ROOT / "applications.json"
SHEET_NAME = "会议支持评估表"
PORT = 8036
HOST = "0.0.0.0"


def as_number(value, default=0):
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def score_monthly_sales(value):
    value = as_number(value, 0)
    if value >= 50:
        return 25
    if value >= 40:
        return 20
    if value >= 30:
        return 15
    if value >= 20:
        return 10
    if value >= 10:
        return 5
    return 1


def get_support(total):
    if total >= 95:
        return {"level": "S++战略项目", "range": "6万+", "amount": 6.5}
    if total >= 90:
        return {"level": "S级重点项目", "range": "6万+", "amount": 6.5}
    if total >= 85:
        return {"level": "A级高价值项目", "range": "5-6万", "amount": 6}
    if total >= 80:
        return {"level": "A级项目", "range": "5-6万", "amount": 5.5}
    if total >= 75:
        return {"level": "B+项目", "range": "4-5万", "amount": 5}
    if total >= 70:
        return {"level": "B级项目", "range": "4-5万", "amount": 4.5}
    if total >= 65:
        return {"level": "B-项目", "range": "3-4万", "amount": 4}
    if total >= 60:
        return {"level": "C+项目", "range": "3-4万", "amount": 3.5}
    if total >= 55:
        return {"level": "C级项目", "range": "2-3万", "amount": 3}
    if total >= 50:
        return {"level": "C-项目", "range": "2-3万", "amount": 2.5}
    if total >= 45:
        return {"level": "D级项目", "range": "1-2万", "amount": 2}
    if total >= 40:
        return {"level": "D-级项目", "range": "1-2万", "amount": 1.5}
    if total >= 30:
        return {"level": "E级项目", "range": "1万", "amount": 1}
    return {"level": "E级项目", "range": "1万以下", "amount": 0.5}


def normalize_record(record):
    form = record.get("form") or {}
    details = record.setdefault("details", {})
    scores = record.setdefault("scores", {})

    details["monthlySales"] = score_monthly_sales(form.get("monthlySales"))
    scores["business"] = (
        as_number(details.get("monthlySales"), 0)
        + as_number(details.get("salesTrend"), 0)
        + as_number(details.get("growthOpportunity"), 0)
    )
    scores["total"] = (
        as_number(scores.get("medical"), 0)
        + as_number(scores.get("strategy"), 0)
        + as_number(scores.get("business"), 0)
        + as_number(scores.get("communication"), 0)
        + as_number(scores.get("execution"), 0)
    )
    record["support"] = get_support(scores["total"])
    return record


def first_empty_application_row(ws):
    for row in range(2, ws.max_row + 2):
        key_values = [ws.cell(row, col).value for col in range(1, 9)]
        if all(value in (None, "") for value in key_values):
            return row
    return ws.max_row + 1


def write_record_to_excel(record):
    form = record.get("form") or {}
    scores = record.get("scores") or {}
    support = record.get("support") or {}

    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]
    ws["AE1"] = "商业价值40分"
    row = first_empty_application_row(ws)

    values = [
        form.get("region", ""),
        form.get("district", ""),
        form.get("applicant", ""),
        form.get("hospital", ""),
        form.get("kol", ""),
        form.get("projectName", ""),
        form.get("meetingDate", ""),
        as_number(form.get("requestAmount"), None),
        form.get("meetingLevel", ""),
        as_number(record.get("details", {}).get("meetingLevel"), None),
        form.get("academicRights", ""),
        as_number(record.get("details", {}).get("academicRights"), None),
        form.get("expertLevel", ""),
        as_number(record.get("details", {}).get("expertLevel"), None),
        form.get("productType", ""),
        as_number(record.get("details", {}).get("productType"), None),
        form.get("hospitalValue", ""),
        as_number(record.get("details", {}).get("hospitalValue"), None),
        as_number(form.get("monthlySales"), None),
        as_number(record.get("details", {}).get("monthlySales"), None),
        form.get("salesTrend", ""),
        as_number(record.get("details", {}).get("salesTrend"), None),
        form.get("growthOpportunity", ""),
        as_number(record.get("details", {}).get("growthOpportunity"), None),
        form.get("communicationValue", ""),
        as_number(record.get("details", {}).get("communicationValue"), None),
        form.get("executionQuality", ""),
        as_number(record.get("details", {}).get("executionQuality"), None),
        as_number(scores.get("medical"), None),
        as_number(scores.get("strategy"), None),
        as_number(scores.get("business"), None),
        as_number(scores.get("communication"), None),
        as_number(scores.get("execution"), None),
        as_number(scores.get("total"), None),
        support.get("range", ""),
        as_number(support.get("amount"), None),
        "",
        record.get("evaluation", ""),
    ]

    for col, value in enumerate(values, start=1):
        ws.cell(row=row, column=col, value=value)

    if form.get("meetingDate"):
        ws.cell(row=row, column=7).number_format = "yyyy-mm-dd"

    wb.save(EXCEL_PATH)
    return row


def read_records():
    if not RECORDS_PATH.exists():
        return []
    try:
        data = json.loads(RECORDS_PATH.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except Exception:
        return []


def write_records(records):
    RECORDS_PATH.write_text(
        json.dumps(records, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def save_record(record, excel_row):
    records = read_records()
    record["excelRow"] = excel_row
    record["serverSavedAt"] = datetime.now().isoformat(timespec="seconds")
    records.insert(0, record)
    write_records(records)
    return records


def response_json(handler, status, body):
    data = json.dumps(body, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(data)))
    handler.end_headers()
    handler.wfile.write(data)


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        super().end_headers()

    def do_GET(self):
        if unquote(self.path).split("?", 1)[0] == "/api/health":
            body = {
                "ok": True,
                "service": "meeting-support-score",
                "excelExists": EXCEL_PATH.exists(),
                "excelFile": str(EXCEL_PATH),
                "recordsFile": str(RECORDS_PATH),
                "recordCount": len(read_records()),
                "time": datetime.now().isoformat(timespec="seconds"),
            }
            response_json(self, 200, body)
            return

        if unquote(self.path).split("?", 1)[0] == "/api/applications":
            response_json(self, 200, {"ok": True, "records": read_records()})
            return

        if self.path == "/":
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self):
        route = unquote(self.path).split("?", 1)[0]
        if route == "/api/clear-applications":
            write_records([])
            response_json(self, 200, {"ok": True, "records": []})
            return

        if route != "/api/save-excel":
            self.send_error(404, "Not found")
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = normalize_record(json.loads(self.rfile.read(length).decode("utf-8")))
            row = write_record_to_excel(payload)
            records = save_record(payload, row)
            body = {
                "ok": True,
                "row": row,
                "file": str(EXCEL_PATH),
                "recordCount": len(records),
                "savedAt": datetime.now().isoformat(timespec="seconds"),
            }
            response_json(self, 200, body)
        except PermissionError as exc:
            body = {
                "ok": False,
                "error": "Excel 文件可能正在打开或被占用，请关闭后重试。",
                "detail": str(exc),
            }
            response_json(self, 423, body)
        except Exception as exc:
            body = {"ok": False, "error": str(exc)}
            response_json(self, 500, body)

    def do_DELETE(self):
        route = unquote(self.path).split("?", 1)[0]
        if route != "/api/applications":
            self.send_error(404, "Not found")
            return

        query = self.path.split("?", 1)[1] if "?" in self.path else ""
        target_id = ""
        for item in query.split("&"):
            key, _, value = item.partition("=")
            if key == "id":
                target_id = unquote(value)
                break

        if not target_id:
            response_json(self, 400, {"ok": False, "error": "Missing id"})
            return

        records = [record for record in read_records() if str(record.get("id")) != target_id]
        write_records(records)
        response_json(self, 200, {"ok": True, "records": records})


if __name__ == "__main__":
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Serving http://0.0.0.0:{PORT}/index.html")
    server.serve_forever()
